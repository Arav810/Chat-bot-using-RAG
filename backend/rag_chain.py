from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain.chains import RetrievalQA
from langchain_groq import ChatGroq
from dotenv import load_dotenv
import os
import fitz
import docx
import pandas as pd
import tempfile
from langchain.schema import Document
from langchain_qdrant import Qdrant
from qdrant_client import QdrantClient
from qdrant_client.http import models as rest
from PIL import Image
import pytesseract
from io import BytesIO

# Load environment variables
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
QDRANT_URL = os.getenv("QDRANT_URL")
QDRANT_API_KEY = os.getenv("QDRANT_API_KEY")

# Embeddings Model
embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")  # Dimension: 384

# Initialize Qdrant Client
qdrant_client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)

try:
    qdrant_client.delete_collection("chatbot-knowledge")
except:
    pass  # Collection may not exist

# Create collection if not exists
if not qdrant_client.collection_exists("chatbot-knowledge"):
    qdrant_client.create_collection(
        collection_name="chatbot-knowledge",
        vectors_config=rest.VectorParams(size=384, distance=rest.Distance.COSINE)
    )

# Vector Store
vector_store = Qdrant(
    client=qdrant_client,
    collection_name="chatbot-knowledge",
    embeddings=embeddings
)

# LLM Setup
llm = ChatGroq(
    groq_api_key=GROQ_API_KEY,
    model="llama3-8b-8192",
    temperature=0.7
)

# Retrieval Chain
retriever = vector_store.as_retriever(search_kwargs={"k": 3})
qa_chain = RetrievalQA.from_chain_type(llm=llm, retriever=retriever)

def run_rag_pipeline(query: str) -> str:
    result = qa_chain.invoke({"query": query})["result"]
    return result

def process_document(file):
    with tempfile.NamedTemporaryFile(delete=False) as tmp:
        tmp.write(file.file.read())
        tmp_path = tmp.name

    text_content = ""
    file_type = ""
    images = []  # store (PIL.Image, page_num)

    # --- PDF ---
    if file.filename.endswith(".pdf"):
        file_type = "pdf"
        with fitz.open(tmp_path) as doc:
            for page_num, page in enumerate(doc, start=1):
                page_text = page.get_text()
                if page_text.strip():
                    text_content += f"[Page {page_num}]\n{page_text}\n"

                for img in page.get_images(full=True):
                    xref = img[0]
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n < 5:
                        img_data = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    else:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                        img_data = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    images.append((img_data, page_num))

    # --- TXT ---
    elif file.filename.endswith(".txt"):
        file_type = "txt"
        with open(tmp_path, "r", encoding="utf-8") as f:
            text_content = f.read()

    # --- DOCX ---
    elif file.filename.endswith(".docx"):
        file_type = "docx"
        document = docx.Document(tmp_path)
        for para in document.paragraphs:
            text_content += para.text + "\n"
        for rel in document.part.rels.values():
            if "image" in rel.target_ref:
                img_bytes = rel.target_part.blob
                img = Image.open(BytesIO(img_bytes))
                images.append((img, None))

    # --- CSV ---
    elif file.filename.endswith(".csv"):
        file_type = "csv"
        df = pd.read_csv(tmp_path)
        text_content = df.to_string(index=False)

    # --- XLSX ---
    elif file.filename.endswith(".xlsx"):
        file_type = "xlsx"
        df = pd.read_excel(tmp_path)
        text_content = df.to_string(index=False)

    # --- Images ---
    elif file.filename.lower().endswith((".png", ".jpg", ".jpeg")):
        file_type = "image"
        img = Image.open(tmp_path)
        images.append((img, None))

    else:
        raise ValueError("Unsupported file format")

    docs = []

    # Normal text chunks
    if text_content.strip():
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
        chunks = text_splitter.split_text(text_content)
        docs.extend([
            Document(
                page_content=chunk,
                metadata={
                    "file_name": file.filename,
                    "file_type": file_type,
                    "source": "text",
                    "page_number": None  # We won't set page here unless PDF per-chunk tracking is added
                }
            )
            for chunk in chunks
        ])

    # OCR chunks from images
    for img, page_num in images:
        ocr_text = pytesseract.image_to_string(img)
        if ocr_text.strip():
            docs.append(Document(
                page_content=ocr_text,
                metadata={
                    "file_name": file.filename,
                    "file_type": file_type,
                    "source": "image",
                    "page_number": page_num
                }
            ))

    # Store in Qdrant
    if docs:
        vector_store.add_documents(docs)
    else:
        raise ValueError("No content extracted from document.")

    return True


# # from langchain.text_splitter import RecursiveCharacterTextSplitter 
# # from langchain_huggingface import HuggingFaceEmbeddings 
# # from langchain.chains import RetrievalQA 
# # from langchain_groq import ChatGroq 
# # from dotenv import load_dotenv 
# # import os 
# # import fitz 
# # import docx 
# # import pandas as pd 
# # import tempfile 
# # from langchain.schema import Document 
# # from langchain_qdrant import Qdrant 
# # from qdrant_client import QdrantClient 
# # from qdrant_client.http import models as rest 
# # # Load environment variables 
# # load_dotenv() 
# # GROQ_API_KEY = os.getenv("GROQ_API_KEY") 
# # QDRANT_URL = os.getenv("QDRANT_URL") 
# # QDRANT_API_KEY = os.getenv("QDRANT_API_KEY") 
# # # Embeddings Model 
# # embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2") # Dimension: 384 
# # # Initialize Qdrant Client 
# # qdrant_client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY) 
# # # Create collection if not exists 
# # if not qdrant_client.collection_exists("chatbot-knowledge"): 
# #     qdrant_client.create_collection( collection_name="chatbot-knowledge", vectors_config=rest.VectorParams(size=384, distance=rest.Distance.COSINE) ) 
# # # Vector Store 
# # vector_store = Qdrant( client=qdrant_client, collection_name="chatbot-knowledge", embeddings=embeddings ) 
# # # LLM Setup 
# # llm = ChatGroq( groq_api_key=GROQ_API_KEY, model="llama3-8b-8192", temperature=0.7 ) 
# # # Retrieval Chain 
# # retriever = vector_store.as_retriever(search_kwargs={"k": 3}) 
# # qa_chain = RetrievalQA.from_chain_type(llm=llm, retriever=retriever) 

# # def run_rag_pipeline(query: str) -> str: 
# #     result = qa_chain.invoke({"query": query})["result"] 
# #     return result 
# # def process_document(file): 
# #     with tempfile.NamedTemporaryFile(delete=False) as tmp: 
# #         tmp.write(file.file.read()) 
# #         tmp_path = tmp.name 
# #     text_content = "" 
# #     file_type = "" 
# #     if file.filename.endswith(".pdf"): 
# #         file_type = "pdf" 
# #         with fitz.open(tmp_path) as doc: 
# #             for page in doc: text_content += page.get_text() 
# #     elif file.filename.endswith(".txt"): 
# #         file_type = "txt" 
# #         with open(tmp_path, "r", encoding="utf-8") as f: 
# #             text_content = f.read() 
# #     elif file.filename.endswith(".docx"): 
# #         file_type = "docx" 
# #         document = docx.Document(tmp_path) 
# #         for para in document.paragraphs: text_content += para.text + "\n" 
# #     elif file.filename.endswith(".csv"): 
# #         file_type = "csv" 
# #         df = pd.read_csv(tmp_path) 
# #         text_content = df.to_string(index=False) 
# #     elif file.filename.endswith(".xlsx"): 
# #         file_type = "xlsx" 
# #         df = pd.read_excel(tmp_path) 
# #         text_content = df.to_string(index=False) 
# #     else: 
# #         raise ValueError("Unsupported file format") 
# #     if not text_content.strip(): 
# #         raise ValueError("Document is empty or text could not be extracted.") 
# # # Split text into chunks 
# #     text_splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50) 
# #     chunks = text_splitter.split_text(text_content) 
# # # Add metadata for file identification 
# #     docs = [ Document(page_content=chunk, metadata={"file_name": file.filename, "file_type": file_type}) 
# #         for chunk in chunks
# #     ] 
# # # Store in Qdrant 
# #     vector_store.add_documents(docs) 

# #     return True